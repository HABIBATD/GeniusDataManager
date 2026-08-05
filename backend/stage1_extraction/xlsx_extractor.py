"""
Stage 1 — XLSX Extractor
=========================
Strategy:
  1. Open with openpyxl (preserves merged cells, avoids formula evaluation issues).
  2. Process each visible worksheet as a separate RawTable.
  3. Detect merged cell regions — expand them so every logical row has all values.
  4. Find the header row: the first row where most cells are non-empty strings.
  5. Preserve XLSX row number for traceability.
  6. Unparseable rows logged, never dropped.
"""

import logging
import re
from typing import Any

from models.intermediate import (
    ExtractionMethod,
    RawCell,
    RawRow,
    RawTable,
    Stage1Result,
)

logger = logging.getLogger(__name__)


def extract_xlsx(file_bytes: bytes, filename: str) -> Stage1Result:
    """
    Extract all visible sheets from an XLSX file.
    Each sheet becomes one RawTable in the Stage1Result.
    """
    try:
        import openpyxl  # type: ignore
        from openpyxl import load_workbook
    except ImportError:
        return Stage1Result(
            filename=filename,
            file_type="xlsx",
            warnings=["openpyxl not installed — XLSX extraction unavailable"],
        )

    import io
    warnings: list[str] = []
    global_unparsed: list[dict] = []
    result_tables: list[RawTable] = []
    table_index = 0

    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=False, data_only=True)
    except Exception as exc:
        return Stage1Result(
            filename=filename,
            file_type="xlsx",
            warnings=[f"Could not open XLSX file: {exc}"],
        )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Skip hidden sheets
        if ws.sheet_state == "hidden":
            warnings.append(f"Sheet '{sheet_name}' is hidden — skipping.")
            continue

        # Expand merged cells: fill each merged region with the top-left value
        merged_map: dict[tuple[int, int], Any] = {}
        for merged_range in ws.merged_cells.ranges:
            top_left_val = ws.cell(merged_range.min_row, merged_range.min_col).value
            for row_idx in range(merged_range.min_row, merged_range.max_row + 1):
                for col_idx in range(merged_range.min_col, merged_range.max_col + 1):
                    merged_map[(row_idx, col_idx)] = top_left_val

        # Read all rows as a 2D list
        all_rows: list[list[Any]] = []
        for xlsx_row in ws.iter_rows():
            row_vals = []
            for cell in xlsx_row:
                pos = (cell.row, cell.column)
                if pos in merged_map:
                    row_vals.append(merged_map[pos])
                else:
                    row_vals.append(cell.value)
            all_rows.append(row_vals)

        if not all_rows:
            warnings.append(f"Sheet '{sheet_name}' is empty — skipping.")
            continue

        # Trim trailing all-None rows and columns
        all_rows = _trim_empty_rows(all_rows)
        if not all_rows:
            continue
        all_rows = _trim_empty_cols(all_rows)

        # Detect header row: first row where ≥50% of cells are non-empty strings
        header_xlsx_row = 0
        for ri, row in enumerate(all_rows):
            non_empty = sum(1 for v in row if v is not None and str(v).strip())
            if non_empty >= max(1, len(row) * 0.5):
                # Check if majority are strings (i.e., looks like a header)
                str_count = sum(1 for v in row if isinstance(v, str))
                if str_count >= non_empty * 0.6:
                    header_xlsx_row = ri
                    break

        raw_header = all_rows[header_xlsx_row]
        headers = [
            str(v).strip() if v is not None and str(v).strip() else f"Col_{ci}"
            for ci, v in enumerate(raw_header)
        ]
        num_cols = len(headers)

        raw_rows: list[RawRow] = []
        row_index = 0
        unparsed: list[dict] = []

        for ri, row in enumerate(all_rows[header_xlsx_row + 1:], start=header_xlsx_row + 1):
            # +1 because openpyxl rows are 1-based; ri is 0-based from all_rows
            xlsx_row_num = ri + 1  # approximate xlsx row number

            # Skip fully blank rows
            if not any(v is not None and str(v).strip() for v in row):
                continue

            # Pad to header width
            padded = list(row) + [None] * max(0, num_cols - len(row))
            padded = padded[:num_cols]

            try:
                cells: list[RawCell] = []
                raw_parts: list[str] = []
                for ci, val in enumerate(padded):
                    raw_text = "" if val is None else str(val).strip()
                    raw_parts.append(raw_text)
                    cells.append(RawCell(
                        col_index=ci,
                        value=_coerce_value(val),
                        raw_text=raw_text,
                    ))
                raw_rows.append(RawRow(
                    row_index=row_index,
                    source_page=None,
                    source_line=xlsx_row_num,
                    raw_text=" | ".join(raw_parts),
                    cells=cells,
                ))
                row_index += 1
            except Exception as exc:
                unparsed.append({
                    "source": f"sheet '{sheet_name}' row {xlsx_row_num}",
                    "reason": str(exc),
                    "raw_text": str(row)[:300],
                })

        global_unparsed.extend(unparsed)
        result_tables.append(RawTable(
            table_index=table_index,
            source_page=None,
            source_sheet=sheet_name,
            extraction_method=ExtractionMethod.OPENPYXL_XLSX,
            headers=headers,
            rows=raw_rows,
            unparsed_rows=unparsed,
        ))
        table_index += 1

    return Stage1Result(
        filename=filename,
        file_type="xlsx",
        total_pages=None,
        tables=result_tables,
        unparsed_rows=global_unparsed,
        warnings=warnings,
    )


def _trim_empty_rows(rows: list[list]) -> list[list]:
    """Remove trailing rows where every cell is None."""
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows


def _trim_empty_cols(rows: list[list]) -> list[list]:
    """Remove trailing columns where every row's cell is None."""
    if not rows:
        return rows
    max_cols = max(len(r) for r in rows)
    last_used_col = max_cols - 1
    for col in range(max_cols - 1, -1, -1):
        if any(col < len(r) and r[col] is not None for r in rows):
            last_used_col = col
            break
    return [r[:last_used_col + 1] for r in rows]


def _coerce_value(val: Any) -> Any:
    """
    Return a typed Python value from a raw openpyxl cell value.
    openpyxl already returns int/float/datetime for typed cells,
    so this mainly handles string cells that look numeric.
    """
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    import datetime
    if isinstance(val, (datetime.datetime, datetime.date)):
        return str(val)
    text = str(val).strip()
    if text == "":
        return None
    cleaned = re.sub(r"[,$£€₹\s]", "", text)
    cleaned = cleaned.replace("(", "-").replace(")", "")
    try:
        f = float(cleaned)
        return int(f) if f == int(f) else f
    except ValueError:
        pass
    return text
