"""
Export Router for GeniusDataManager.
GET /api/export/{result_id}?format=csv|xlsx|json
"""
import io
import json
import logging
from typing import Optional
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment as XLAlignment
import pandas as pd

from services.storage import storage

logger = logging.getLogger("genius.router.export")
router = APIRouter(prefix="/api", tags=["Export"])

ALLOWED_FORMATS = {"csv", "xlsx", "json"}


@router.get("/export/{result_id}")
async def export_result(
    result_id: str,
    format: str = Query("csv", description="'csv' | 'xlsx' | 'json'"),
):
    """
    Export the analysis results or aligned data as CSV, XLSX, or JSON file attachment.
    """
    fmt = format.lower()
    if fmt not in ALLOWED_FORMATS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported export format '{fmt}'. Choose from: csv, xlsx, json.",
        )

    item = storage.get_analysis_result(result_id)
    # Check if result_id is an aligned_id directly
    aligned_df = None
    analysis_res = None

    if item:
        analysis_res, aligned_df = item
    else:
        # Check if result_id is directly an aligned_id in dataframes
        df_direct = storage.get_dataframe(result_id)
        if df_direct is not None:
            aligned_df = df_direct
        else:
            raise HTTPException(status_code=404, detail=f"Result ID or Aligned ID '{result_id}' not found.")

    filename_base = f"genius_export_{result_id[:8]}"

    # 1. CSV Export
    if fmt == "csv":
        buffer = io.StringIO()
        if aligned_df is not None and not aligned_df.empty:
            aligned_df.to_csv(buffer, index=False)
        elif analysis_res and analysis_res.table_data:
            pd.DataFrame(analysis_res.table_data).to_csv(buffer, index=False)
        else:
            buffer.write("No tabular data available for export.\n")

        content_bytes = buffer.getvalue().encode("utf-8")
        return StreamingResponse(
            io.BytesIO(content_bytes),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'},
        )

    # 2. JSON Export
    elif fmt == "json":
        export_payload = {}
        if analysis_res:
            export_payload["analysis"] = analysis_res.model_dump()
        if aligned_df is not None:
            export_payload["data"] = aligned_df.to_dict(orient="records")

        content_bytes = json.dumps(export_payload, indent=2, default=str).encode("utf-8")
        return StreamingResponse(
            io.BytesIO(content_bytes),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.json"'},
        )

    # 3. XLSX Export
    elif fmt == "xlsx":
        wb = openpyxl.Workbook()
        # Sheet 1: Executive Summary
        ws_summary = wb.active
        ws_summary.title = "Executive Summary"

        header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        title_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")

        ws_summary["A1"] = "GeniusDataManager — Executive Summary"
        ws_summary["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        ws_summary["A1"].fill = title_fill

        row_idx = 3
        if analysis_res:
            ws_summary.cell(row=row_idx, column=1, value="Analysis Type:").font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2, value=analysis_res.analysis_type.title())
            row_idx += 1

            ws_summary.cell(row=row_idx, column=1, value="Summary:").font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2, value=analysis_res.summary)
            row_idx += 2

            if analysis_res.kpis:
                ws_summary.cell(row=row_idx, column=1, value="Key Performance Indicators:").font = Font(size=12, bold=True)
                row_idx += 1
                for kpi in analysis_res.kpis:
                    ws_summary.cell(row=row_idx, column=1, value=kpi.label).font = Font(bold=True)
                    ws_summary.cell(row=row_idx, column=2, value=kpi.formatted)
                    if kpi.subtext:
                        ws_summary.cell(row=row_idx, column=3, value=kpi.subtext)
                    row_idx += 1
                row_idx += 1

        # Sheet 2: Data Records
        ws_data = wb.create_sheet(title="Aligned Data")
        export_df = aligned_df if (aligned_df is not None and not aligned_df.empty) else (
            pd.DataFrame(analysis_res.table_data) if (analysis_res and analysis_res.table_data) else pd.DataFrame()
        )

        if not export_df.empty:
            # Write headers
            for col_idx, col_name in enumerate(export_df.columns, start=1):
                cell = ws_data.cell(row=1, column=col_idx, value=str(col_name))
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.fill = header_fill

            # Write rows
            for r_idx, (_, row) in enumerate(export_df.iterrows(), start=2):
                for c_idx, val in enumerate(row, start=1):
                    ws_data.cell(row=r_idx, column=c_idx, value=str(val) if pd.notna(val) else "")

        output_buffer = io.BytesIO()
        wb.save(output_buffer)
        output_buffer.seek(0)

        return StreamingResponse(
            output_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'},
        )
